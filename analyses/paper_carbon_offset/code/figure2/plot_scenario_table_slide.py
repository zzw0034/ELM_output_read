"""
Talk slide: the three land-management scenarios in one table (user request 2026-10-07), as a 16:9 PNG and as an
editable one-slide PowerPoint table. Definitions as locked on 2026-08-19 (notes / scenario memory):
  RF  restore forest on grassland where forest existed in 1850 (+142 x 10^3 km2, ramped 2024-2050) and set all
      harvest to zero (restoration plus a region-wide harvest ban); one land-use file shared by all SSPs
  RH  annual wood-harvest rate halved (rotation doubled); land cover as Default
  DF  all forest converted to grass in 2024 and kept clear: a counterfactual, used as Default - DF
Colour chips = the Figure 2 scenario colours (plot_scenario_management_benefits.COLORS).

Runs locally, from the analysis root (the .pptx needs python-pptx, e.g. via uv):
    uv run --no-project --with python-pptx --with matplotlib python code/figure2/plot_scenario_table_slide.py
Output: figures/figure2/slides/management_scenarios_table.{png,pptx}
"""
import os
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from plot_scenario_management_benefits import COLORS, GRID, INK, INK2  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(ROOT, "figures/figure2/slides")
HEAD = ["Scenario", "What it does", "Benefit measured as"]
ROWS = [
    ("RF", "Restoration\nand protection",
     "Restore forest on grassland that\nwas forest in 1850 (+142 × 10³ km²,\nphased in 2024–2050); stop all harvest",
     "RF − Default"),
    ("RH", "Reduced\nharvest",
     "Annual wood-harvest rate halved\n(rotation doubled); forest area\nsame as Default",
     "RH − Default"),
    ("DF", "Deforestation\n(counterfactual)",
     "All forest converted to grassland\nin 2024 and kept clear to 2100",
     "Default − DF\n(value of protecting\nexisting forest)"),
]
COLX = [0.04, 0.37, 0.76]          # left edges of the three columns (figure fraction)


def png():
    fig = plt.figure(figsize=(16, 9), facecolor="white")
    fig.text(0.04, 0.92, "Three land-management scenarios, each run under all four SSPs", fontsize=30, color=INK,
             fontweight="bold", va="center")
    y = 0.80
    for x, h in zip(COLX, HEAD):
        fig.text(x, y, h, fontsize=22, color=INK2, fontweight="bold", va="center")
    fig.add_artist(plt.Line2D([0.04, 0.96], [y - 0.045, y - 0.045], color=INK, lw=1.5, transform=fig.transFigure))
    row_h = 0.205
    for i, (code, name, what, comp) in enumerate(ROWS):
        yc = y - 0.045 - row_h * (i + 0.5)
        fig.add_artist(FancyBboxPatch((COLX[0], yc - 0.035), 0.065, 0.07, boxstyle="round,pad=0.004,rounding_size=0.012",
                                      transform=fig.transFigure, facecolor=COLORS[code], edgecolor="none"))
        fig.text(COLX[0] + 0.0325, yc, code, fontsize=26, color="white", fontweight="bold", ha="center", va="center")
        fig.text(COLX[0] + 0.08, yc, name, fontsize=22, color=INK, fontweight="bold", va="center",
                 linespacing=1.3)
        fig.text(COLX[1], yc, what, fontsize=19, color=INK, va="center", linespacing=1.4)
        fig.text(COLX[2], yc, comp, fontsize=19, color=INK, va="center", linespacing=1.4,
                 fontweight="bold" if "\n" not in comp else "normal")
        if i < len(ROWS) - 1:
            yl = y - 0.045 - row_h * (i + 1)
            fig.add_artist(plt.Line2D([0.04, 0.96], [yl, yl], color=GRID, lw=1.5, transform=fig.transFigure))
    fig.text(0.04, 0.06, "Default follows each SSP's own land use and harvest. RF uses the same land-use path in every SSP.", fontsize=16, color=INK2, va="center")
    os.makedirs(OUT, exist_ok=True)
    p = os.path.join(OUT, "management_scenarios_table.png")
    fig.savefig(p, dpi=200, facecolor="white")
    plt.close(fig)
    return p


def pptx():
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    rgb = lambda h: RGBColor.from_string(h.lstrip("#").upper())
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[5])           # title only
    s.shapes.title.text = "Three land-management scenarios, each run under all four SSPs"
    for r in s.shapes.title.text_frame.paragraphs[0].runs:
        r.font.size, r.font.bold, r.font.name = Pt(32), True, "Calibri"
    s.shapes.title.left, s.shapes.title.top, s.shapes.title.width, s.shapes.title.height = Inches(0.5), Inches(0.3), Inches(12.3), Inches(1.0)
    rows, cols = len(ROWS) + 1, 3
    tbl = s.shapes.add_table(rows, cols, Inches(0.5), Inches(1.6), Inches(12.3), Inches(5.0)).table
    for j, w in enumerate((3.3, 5.6, 3.4)):
        tbl.columns[j].width = Inches(w)
    def cell(i, j, text, size=18, bold=False, color=INK, fill=None):
        c = tbl.cell(i, j)
        c.text = ""
        tf = c.text_frame
        tf.word_wrap = True
        for k, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = line
            run.font.size, run.font.bold, run.font.name = Pt(size), bold, "Calibri"
            run.font.color.rgb = rgb(color)
        c.fill.solid()
        c.fill.fore_color.rgb = rgb(fill or "#FFFFFF")
    for j, h in enumerate(HEAD):
        cell(0, j, h, 18, True, "#FFFFFF", INK2)
    for i, (code, name, what, comp) in enumerate(ROWS, 1):
        cell(i, 0, f"{code}  " + name.replace("\n", " "), 18, True, "#FFFFFF", COLORS[code])
        cell(i, 1, what.replace("\n", " "), 16)
        cell(i, 2, comp, 16, bold="\n" not in comp)
    p = os.path.join(OUT, "management_scenarios_table.pptx")
    prs.save(p)
    return p


if __name__ == "__main__":
    print("wrote", png())
    try:
        print("wrote", pptx())
    except ImportError:
        print("python-pptx not available: PNG only")
